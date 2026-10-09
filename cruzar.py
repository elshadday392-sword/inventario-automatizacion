import csv
import openpyxl
from docx import Document

# --- 1. Leer el Word y sacar estados por producto ---
doc = Document("reporte.docx")
estados = {}

for tabla in doc.tables:
    for fila in tabla.rows[1:]:  # [1:] salta el encabezado
        producto = fila.cells[0].text.strip()
        estado = fila.cells[2].text.strip()
        estados[producto] = estado

print("Estados leídos del Word:")
for k, v in estados.items():
    print(f"  {k} -> {v}")
print()

# --- 2. Leer el Excel y cruzar con los estados ---
wb = openpyxl.load_workbook("ventas.xlsx")
ws = wb.active

reporte = []
for fila in ws.iter_rows(min_row=2, values_only=True):  # salta encabezado
    producto, stock, vendidos, precio = fila
    estado = estados.get(producto, "SIN ESTADO")
    reporte.append({
        "producto": producto,
        "stock": stock,
        "vendidos": vendidos,
        "precio": precio,
        "estado": estado
    })

# --- 3. Imprimir el reporte combinado ---
print("=== REPORTE COMBINADO ===")
for r in reporte:
    print(f"{r['producto']:<18} | stock: {r['stock']:<3} | precio: {r['precio']:<6} | estado: {r['estado']}")

# --- 4. Guardar como CSV ---
with open("reporte_final.csv", "w", newline="", encoding="utf-8") as archivo:
    campos = ["producto", "stock", "vendidos", "precio", "estado"]
    escritor = csv.DictWriter(archivo, fieldnames=campos)
    escritor.writeheader()
    escritor.writerows(reporte)

print(f"\nTotal: {len(reporte)} productos guardados en reporte_final.csv")
