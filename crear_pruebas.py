from openpyxl import Workbook
from docx import Document

# --- Excel ---
wb = Workbook()
ws = wb.active
ws.title = "Ventas"
ws.append(["producto", "stock", "vendidos", "precio"])
ws.append(["Filtro de aceite", 45, 3, 8.5])
ws.append(["Pastilla de freno", 8, 12, 25.0])
ws.append(["Bujía NGK", 3, 15, 4.75])
ws.append(["Batería 12V", 2, 11, 85.0])
wb.save("ventas.xlsx")
print("ventas.xlsx creado")

# --- Word ---
doc = Document()
doc.add_heading("Reporte de inventario", 0)
doc.add_paragraph("Este es un reporte de prueba.")

tabla = doc.add_table(rows=3, cols=3)
tabla.style = "Table Grid"

encabezados = ["Producto", "Stock", "Estado"]
for i, texto in enumerate(encabezados):
    tabla.rows[0].cells[i].text = texto

tabla.rows[1].cells[0].text = "Bujía NGK"
tabla.rows[1].cells[1].text = "3"
tabla.rows[1].cells[2].text = "CRÍTICO"

tabla.rows[2].cells[0].text = "Batería 12V"
tabla.rows[2].cells[1].text = "2"
tabla.rows[2].cells[2].text = "CRÍTICO"

doc.save("reporte.docx")
print("reporte.docx creado")
