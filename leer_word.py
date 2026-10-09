from docx import Document

doc = Document("reporte.docx")

print("=== PÁRRAFOS DEL WORD ===")
for p in doc.paragraphs:
    if p.text.strip():
        print(p.text)

print("\n=== TABLAS DEL WORD ===")
for tabla in doc.tables:
    for fila in tabla.rows:
        celdas = [celda.text for celda in fila.cells]
        print(" | ".join(celdas))
